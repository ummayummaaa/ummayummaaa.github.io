from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


OUTPUT = Path(__file__).resolve().parents[1] / "student-test-example.docx"


def set_font(run, name="Arial", size=11, bold=None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:cs"), name)
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold


def add_line(document, text, bold=False, space_after=4):
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(space_after)
    run = paragraph.add_run(text)
    set_font(run, bold=bold)
    return paragraph


def remove_paragraph_borders(paragraph_or_style):
    paragraph_properties = paragraph_or_style._element.get_or_add_pPr()
    borders = paragraph_properties.find(qn("w:pBdr"))
    if borders is not None:
        paragraph_properties.remove(borders)


def main():
    document = Document()
    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)

    styles = document.styles
    for style_name in ("Normal", "Title"):
        style = styles[style_name]
        style.font.name = "Arial"
        style._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Arial")
        style._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Arial")
        style.font.color.rgb = None
    styles["Normal"].font.size = Pt(11)
    styles["Title"].font.size = Pt(20)
    styles["Title"].font.bold = True
    remove_paragraph_borders(styles["Title"])

    title = document.add_paragraph(style="Title")
    title.paragraph_format.space_after = Pt(10)
    remove_paragraph_borders(title)
    title_run = title.add_run("Пример оформления пользовательского теста")
    set_font(title_run, size=20, bold=True)

    intro = document.add_paragraph()
    intro.paragraph_format.space_after = Pt(14)
    intro_run = intro.add_run(
        "Каждый вопрос начинается с номера. Варианты начинаются с A), B), C). "
        "Полностью выделите жирным ровно один правильный вариант вместе с буквой."
    )
    set_font(intro_run)

    add_line(document, "1. Какой вариант является правильным?", space_after=7)
    add_line(document, "A) Первый вариант")
    add_line(document, "B) Второй вариант", bold=True)
    add_line(document, "C) Третий вариант", space_after=7)
    add_line(
        document,
        "Объяснение: Второй вариант правильный, потому что он полностью выделен жирным.",
    )
    add_line(
        document,
        "Дополнительный абзац объяснения также будет прикреплён к первому вопросу.",
    )
    add_line(document, "Источник: Название учебника, глава 5.", space_after=13)

    add_line(document, "2. Какой орган перекачивает кровь?", space_after=7)
    add_line(document, "A) Сердце", bold=True)
    add_line(document, "B) Лёгкие")
    add_line(document, "C) Печень")

    core_properties = document.core_properties
    core_properties.title = "Пример пользовательского теста MedQuiz"
    core_properties.subject = "Шаблон DOCX для импорта вопросов"
    core_properties.author = "MedQuiz"
    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
